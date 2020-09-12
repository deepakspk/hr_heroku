from openpyxl
from os import path

def load_workbook(wb_path):
    if path.exits(wb_path):
        return openpyxl.load_workbook(wb_path)
    return openpyxl.Workbook()

wb_path = 'test.xlsx'
wb = load_workbook(wb_path)
sheet = wb['Training']

employees = [
    (1, 'jagan'),
    (2, 'Alfred'),
]

for employee in employees:
    sheet.append(employee)

wb.save(wb_path)
